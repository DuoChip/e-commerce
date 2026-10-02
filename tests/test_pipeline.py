"""Checks for the core accounting and star-schema invariants."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from data_cleaning import clean_transactions  # noqa: E402
from load_sqlserver import build_load_sql, inspect_exports, prepare_csv_for_sql  # noqa: E402
from rfm_segmentation import calculate_rfm  # noqa: E402
from star_schema import build_star_schema, export_star_schema  # noqa: E402


class PipelineTest(unittest.TestCase):
    def test_original_dates_returns_and_foreign_keys(self):
        raw = pd.DataFrame([
            ["100", "A", "Product A", 2, "12/1/2010 8:26", 10.0, 1, "UK"],
            ["C101", "A", "Product A", -1, "12/2/2010 9:00", 10.0, 1, "UK"],
            ["200", "B", "Product B", 1, "12/3/2010 9:00", 5.0, 2, "France"],
            ["C300", "B", "Product B", -1, "12/4/2010 9:00", 5.0, 3, "France"],
        ], columns=["InvoiceNo", "StockCode", "Description", "Quantity",
                    "InvoiceDate", "UnitPrice", "CustomerID", "Country"])
        clean = clean_transactions(raw)
        self.assertEqual(str(clean["InvoiceDate"].min().date()), "2010-12-01")
        rfm = calculate_rfm(clean)
        self.assertEqual(rfm.set_index("CustomerID").loc[1, "Frequency"], 1)
        self.assertEqual(rfm.set_index("CustomerID").loc[1, "Monetary"], 10)
        self.assertEqual(rfm.set_index("CustomerID").loc[3, "Status"], "No Purchase")
        self.assertEqual(rfm.set_index("CustomerID").loc[1, "F_Score"],
                         rfm.set_index("CustomerID").loc[2, "F_Score"])

        tables = build_star_schema(clean, rfm)
        fact = tables["fact_transactions"]
        self.assertEqual(len(fact), 4)
        self.assertEqual(fact["gross_sales"].sum(), 25)
        self.assertEqual(fact["return_amount"].sum(), 15)
        self.assertEqual(fact["net_sales"].sum(), 10)
        with tempfile.TemporaryDirectory() as temp:
            workbook_path = export_star_schema(tables, Path(temp))
            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
            self.assertEqual(set(workbook.sheetnames), set(tables))
            fact_rows = list(workbook["fact_transactions"].values)
            self.assertEqual(len(fact_rows), 5)  # header plus four transactions
            self.assertEqual(fact_rows[1][3], 20101201)
            self.assertEqual(fact_rows[-1][3], 20101204)
            workbook.close()
            counts, net_sales = inspect_exports(Path(temp))
            self.assertEqual(counts["fact_transactions"], 4)
            self.assertEqual(float(net_sales), 10)
            sql = build_load_sql(counts, net_sales, "/tmp/ecommerce_test", False)
            self.assertLess(sql.index("BULK INSERT [ecommerce].[dim_date]"),
                            sql.index("BULK INSERT [ecommerce].[fact_transactions]"))
            self.assertNotIn("DELETE FROM", sql)

    def test_sql_csv_uses_fixed_point_decimal(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "raw.csv"
            destination = Path(temp) / "prepared.csv"
            source.write_text("customer_key,monetary_net\n1,5.684341886080802e-14\n")
            prepare_csv_for_sql("dim_customer", source, destination)
            self.assertIn("1,0.000000", destination.read_text())


if __name__ == "__main__":
    unittest.main()
